from __future__ import annotations

import json
import os
from datetime import datetime
from email.message import EmailMessage
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import smtplib

from openpyxl import Workbook, load_workbook

DATA_DIR = Path(__file__).parent / "data"
DATA_DIR.mkdir(exist_ok=True)
WORKBOOK_PATH = DATA_DIR / "contacts.xlsx"

SMTP_HOST = os.getenv("SMTP_HOST")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER")
SMTP_PASS = os.getenv("SMTP_PASS")
SMTP_TO = os.getenv("SMTP_TO", "nroy9520@gmail.com")


def load_workbook_file() -> Workbook:
    if WORKBOOK_PATH.exists():
        return load_workbook(WORKBOOK_PATH)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contacts"
    sheet.append(["Timestamp", "Name", "Email", "Company", "Message"])
    workbook.save(WORKBOOK_PATH)
    return workbook


def append_contact(data: dict[str, str]) -> None:
    workbook = load_workbook_file()
    sheet = workbook.active
    sheet.append(
        [
            datetime.utcnow().isoformat(timespec="seconds"),
            data.get("name", ""),
            data.get("email", ""),
            data.get("company", ""),
            data.get("message", ""),
        ]
    )
    workbook.save(WORKBOOK_PATH)


def send_report() -> bool:
    if not all([SMTP_HOST, SMTP_USER, SMTP_PASS, SMTP_TO]):
        return False

    message = EmailMessage()
    message["Subject"] = "Origon AI contact requests"
    message["From"] = SMTP_USER
    message["To"] = SMTP_TO
    message.set_content("Attached is the latest contact request export.")

    if WORKBOOK_PATH.exists():
        message.add_attachment(
            WORKBOOK_PATH.read_bytes(),
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=WORKBOOK_PATH.name,
        )

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASS)
        server.send_message(message)
    return True


class ContactHandler(SimpleHTTPRequestHandler):
    def do_POST(self) -> None:
        if self.path != "/contact":
            self.send_error(404, "Not found")
            return

        content_length = int(self.headers.get("Content-Length", "0"))
        payload = self.rfile.read(content_length).decode("utf-8")

        try:
            data = json.loads(payload)
        except json.JSONDecodeError:
            self.send_error(400, "Invalid JSON")
            return

        append_contact(data)
        send_report()
        self.send_response(204)
        self.end_headers()


if __name__ == "__main__":
    server = ThreadingHTTPServer(("0.0.0.0", 8000), ContactHandler)
    print("Serving on http://0.0.0.0:8000")
    server.serve_forever()
